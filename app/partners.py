"""Past collaborations imported from the company's own tracker (Excel or CSV).

Used to flag creators the company already worked with, to fill Agency and Year-week in downloads,
and to show the AI what kind of creators have worked for this brand before.
"""
import csv
import io
import re

from .markets import MARKETS

# Our field -> header names we accept (lowercase). Prenew's sheet uses the first name in each list.
HEADERS = {
    "key": ["creator key", "creator", "influencer", "name", "creator name"],
    "channel": ["creator / channel", "channel", "handle", "account"],
    "market": ["market"],
    "country": ["country"],
    "agency": ["agency"],
    "week": ["year-week", "week", "date", "campaign week"],
    "platform": ["platform", "platforms"],
    "niche": ["niche / content", "niche", "content"],
    "yt_subs": ["yt subscribers", "youtube subscribers", "subscribers"],
    "tt_followers": ["tiktok followers"],
}
YES = {"yes", "y", "x", "true", "1", "kyllä", "ja"}
COUNTRY_CODES = {v["name"].lower(): k for k, v in MARKETS.items()}


class TrackerError(ValueError):
    pass


def norm(text) -> str:
    return re.sub(r"[^a-z0-9]", "", str(text or "").lower())


def _rows_xlsx(data: bytes) -> list[list]:
    from openpyxl import load_workbook
    try:
        wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    except Exception as e:
        raise TrackerError("Couldn't read that Excel file") from e
    ws = wb.worksheets[0]
    return [list(r) for r in ws.iter_rows(values_only=True)]


def _rows_csv(data: bytes) -> list[list]:
    for encoding in ("utf-8-sig", "cp1252"):
        try:
            text = data.decode(encoding)
            break
        except UnicodeDecodeError:
            continue
    else:
        raise TrackerError("Couldn't read that CSV file")
    dialect = csv.Sniffer().sniff(text[:4000], delimiters=",;\t") if text.strip() else csv.excel
    return [row for row in csv.reader(io.StringIO(text), dialect)]


def _header_map(row: list) -> dict[str, int]:
    names = [str(v or "").strip().lower() for v in row]
    found = {}
    for field, options in HEADERS.items():
        for option in options:
            if option in names:
                found[field] = names.index(option)
                break
    return found


def _read(data: bytes, filename: str) -> list[list]:
    return _rows_xlsx(data) if filename.lower().endswith((".xlsx", ".xlsm")) else _rows_csv(data)


def sheet(data: bytes, filename: str) -> dict:
    """The tracker as it was uploaded (header + non-empty rows), so it can be handed back completed."""
    rows = _read(data, filename)
    for i, row in enumerate(rows[:10]):
        cols = _header_map(row)
        if "key" in cols or "channel" in cols:
            width = max(j for j, v in enumerate(row) if v not in (None, "")) + 1
            body = [[_cell(v) for v in (list(r) + [None] * width)[:width]] for r in rows[i + 1:]]
            return {"header": [str(v or "").strip() for v in row[:width]],
                    "rows": [r for r in body if any(v not in (None, "") for v in r)]}
    raise TrackerError("No 'Creator' or 'Creator / channel' column found in the first rows")


def _cell(v):
    """Numbers stay numbers (whole ones as int), everything else becomes text; JSON-safe."""
    if isinstance(v, bool) or v is None:
        return v
    if isinstance(v, (int, float)):
        return int(v) if float(v).is_integer() else v
    return str(v).strip() or None


def parse(data: bytes, filename: str) -> list[dict]:
    """Rows of a collaboration tracker -> one entry per creator (repeat collaborations merged)."""
    rows = _read(data, filename)
    start, cols = None, {}
    for i, row in enumerate(rows[:10]):
        cols = _header_map(row)
        if "key" in cols or "channel" in cols:
            start = i + 1
            break
    if start is None:
        raise TrackerError("No 'Creator' or 'Creator / channel' column found in the first rows")

    def get(row, field):
        i = cols.get(field)
        v = row[i] if i is not None and i < len(row) else None
        return str(v).strip() if v not in (None, "") else ""

    people: dict[str, dict] = {}
    for row in rows[start:]:
        name = get(row, "key") or get(row, "channel")
        if not name:
            continue
        p = people.setdefault(norm(name), {"name": name, "aliases": [], "channels": [], "weeks": [], "platforms": [],
                                           "market": "", "niche": "", "agency": False, "collabs": 0,
                                           "yt_subs": None, "tt_followers": None})
        p["collabs"] += 1
        for raw in (name, get(row, "channel")):
            if raw and raw not in p["channels"]:
                p["channels"].append(raw)
        for alias in _aliases(name, get(row, "channel")):
            if alias not in p["aliases"]:
                p["aliases"].append(alias)
        market = get(row, "market").upper()[:2] or COUNTRY_CODES.get(get(row, "country").lower(), "")
        p["market"] = market or p["market"]
        p["agency"] = p["agency"] or get(row, "agency").lower() in YES
        week = get(row, "week")
        if week and week not in p["weeks"]:
            p["weeks"].append(week)
        platform = get(row, "platform")
        if platform and platform not in p["platforms"]:
            p["platforms"].append(platform)
        p["niche"] = get(row, "niche") or p["niche"]
        for field in ("yt_subs", "tt_followers"):
            try:
                p[field] = int(float(get(row, field))) if get(row, field) else p[field]
            except ValueError:
                pass
    for p in people.values():
        p["weeks"].sort()
    if not people:
        raise TrackerError("The file has a header row but no creators under it")
    return list(people.values())


def _aliases(*names) -> list[str]:
    """"Alex (AlexMedia)" -> alexalexmedia, alex, alexmedia; "Kim + Robin" -> kimrobin, kim, robin."""
    out = []
    for name in names:
        for part in [name, *re.split(r"[+&/,()]| and ", name)]:
            a = norm(part)
            if len(a) >= 3 and a not in out:
                out.append(a)
    return out


def index(company: dict) -> dict[str, dict]:
    return {a: p for p in (company.get("partners") or {}).get("items", []) for a in p["aliases"]}


def find(idx: dict[str, dict], creator: dict, match: dict | None = None) -> dict | None:
    """The past collaboration this creator is, if any. Short names (e.g. "Kim") also need the same market."""
    if not idx:
        return None
    country = (match or {}).get("country") or creator.get("country") or ""
    markets = {country, *((match or {}).get("search_markets") or [])} - {""}
    for key in (norm(creator.get("name")), norm(creator.get("handle"))):
        p = idx.get(key)
        if p and (len(key) >= 6 or p["market"] in markets):
            return p
    return None


def brief(company: dict, limit: int = 30) -> str:
    """Past partners in a few tokens, for the AI: what has worked for this brand."""
    items = (company.get("partners") or {}).get("items", [])
    if not items:
        return ""
    items = sorted(items, key=lambda p: (p["collabs"], p["weeks"][-1] if p["weeks"] else ""), reverse=True)[:limit]

    def size(p):
        parts = [f"YT {_short(p['yt_subs'])}" if p["yt_subs"] else "", f"TikTok {_short(p['tt_followers'])}" if p["tt_followers"] else ""]
        return ", ".join(x for x in parts if x)

    lines = [", ".join(x for x in (p["market"], " / ".join(p["platforms"][:1]), p["niche"], size(p)) if x) for p in items]
    return "; ".join(f"{p['name']} ({line})" for p, line in zip(items, lines))


def _short(n: int) -> str:
    return f"{n / 1e6:.1f}M" if n >= 1e6 else f"{n / 1e3:.0f}k" if n >= 1e3 else str(n)
