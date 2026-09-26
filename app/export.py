"""Excel/CSV downloads in the layout of Prenew's own collaboration tracker.

The first 12 columns match their sheet exactly (one row per creator, YouTube and TikTok side by side),
so rows can be pasted straight into it. Scout's extra details (contacts, links, score, risks) follow.
"""
import csv
import io
import re

from . import linking, partners
from .markets import MARKETS, PLATFORMS
from .metrics import agency_hint

TRACKER_COLUMNS = [
    ("Creator key", 18), ("Market", 8), ("Country", 11), ("Creator / channel", 26), ("Agency", 8),
    ("Year-week", 10), ("Platform", 17), ("Niche / content", 22), ("YT subscribers", 12),
    ("YT views / video", 12), ("TikTok followers", 12), ("TikTok views / video", 12),
]
EXTRA_COLUMNS = [
    ("Instagram followers", 12), ("Email", 30), ("YouTube", 34), ("TikTok", 34), ("Instagram", 34),
    ("Other links", 40), ("Match score", 9), ("Fit", 7), ("Audience quality", 9), ("Confidence", 11),
    ("Authenticity", 10), ("Est. price per post (EUR)", 14), ("Views counted over", 20), ("Views trend", 10),
    ("Engagement vs typical", 12), ("Last post (days ago)", 10), ("Risks / red flags", 44),
    ("Why they fit", 60), ("Verdict", 60), ("Summary", 60), ("Status", 11), ("Past collaborations", 22),
]
COLUMNS = TRACKER_COLUMNS + EXTRA_COLUMNS
NUMBER_COLS = {"YT subscribers", "TikTok followers", "Instagram followers"}
VIEWS_COLS = {"YT views / video", "TikTok views / video"}
URL_COLS = {"YouTube", "TikTok", "Instagram"}
# Shows 40213 as "40K avg" like Prenew's sheet, while the cell stays a number you can sort and sum.
VIEWS_FORMAT = '[>=1000000]0.0,,"M avg";[>=1000]0,"K avg";0" avg"'

ABOUT = [
    ("Columns A-L", "Same layout as Prenew's collaboration tracker, one row per creator. "
                    "A creator's YouTube and TikTok are merged into one row when one profile links to the other."),
    ("Market / Country", "Where the creator and their audience are, judged by the AI from language, content and profile."),
    ("Agency", "Yes = likely reached through an agency or management (a company email address that isn't "
               "the creator's own, or management mentioned in the bio). A guess: check before outreach. "
               "For past partners, taken from your tracker."),
    ("Year-week", "Your latest collaboration with this creator (from your imported tracker). Empty = new creator."),
    ("Views / video", "Average views per video from the last 30 days (90 days for less active creators; "
                      "posts under 2 days old left out). YouTube: normal videos, Shorts left out when there are enough."),
    ("Views trend", "Average views in the last 30 days compared with the 60 days before."),
    ("Engagement vs typical", "Likes + comments per view, and views per follower, compared with typical accounts "
                              "of the same size on the same platform. 1.0 = typical, 2.0 = twice as engaged."),
    ("Match score", "0-100, used for ranking: a blend of Fit and Audience quality. The brand's campaign goal decides "
                    "the blend (Sales 65/35, Balanced 60/40, Awareness 50/50)."),
    ("Fit", "0-100: would a marketer pick them for this brand? Content (matches the creator types), audience "
            "(viewers are the brand's customers: age, interests, trust), market (audience in the target markets, "
            "incl. comment language), brand & safety, and readiness & cost (contact, sponsor experience, price vs budget). "
            "Every part is backed by evidence that cites posts or comments."),
    ("Audience quality", "0-100: is the audience real and paying attention? Authenticity 35%, engagement vs typical "
                         "30%, consistency 15% (how steady views are across the middle half of posts), activity 10%, "
                         "momentum 10% (views trend)."),
    ("Confidence", "How much the scores rest on: whether the AI read their posts (or did a deep evaluation), "
                   "how many posts, likes and comments were available."),
    ("Authenticity", "Starts at a neutral 70 (nothing suspicious found). Raised by positive evidence (many followers "
                     "watching, real conversation in comments), lowered by signals such as few followers watching, likes "
                     "far above typical, many likes but no comments, or generic/copy-paste comments. Capped when there is "
                     "little data. Signals, not proof."),
    ("Est. price per post", "A rough range: median views × common rates per 1,000 views (YouTube €15-30, TikTok €8-18). "
                            "Check with the creator."),
]


def _short_name(name: str) -> str:
    """"Julian Ivanov | KI-Automatisierung" -> "Julian Ivanov" (but "NE | Tech" stays whole)."""
    first = re.split(r"\s[|•·–-]\s", name or "")[0].strip()
    return first if len(first) >= 4 else (name or "").strip()


def _niche(m: dict) -> str:
    games = " / ".join(m.get("games", [])[:2])
    niche = m.get("niche") or (m.get("tags") or [""])[0]
    if games and (not niche or re.search(r"gam|play", niche, re.I)):
        return games
    return niche


def _platform_label(c: dict) -> str:
    if c["platform"] == "youtube" and (c.get("shorts_share") or 0) >= 0.7:
        return "YouTube Shorts"
    return PLATFORMS[c["platform"]]


def _row(primary: dict, m: dict, profs: dict[str, dict], partner: dict | None, extra_matches: list[dict]) -> dict:
    yt, tt, ig = profs.get("youtube"), profs.get("tiktok"), profs.get("instagram")
    ordered = [primary] + [p for p in profs.values() if p is not primary]
    country = m.get("country") or primary.get("country") or next((p.get("country") for p in ordered if p.get("country")), "")
    if not country and len(m.get("search_markets") or []) == 1:
        country = m["search_markets"][0]
    names = []
    for p in ordered:
        if partners.norm(p.get("name")) not in {partners.norm(n) for n in names}:
            names.append(p.get("name") or "")
    socials = {}
    for p in ordered:
        for net, url in (p.get("socials") or {}).items():
            socials.setdefault(net, url)
    emails = list(dict.fromkeys(e for p in ordered for e in p.get("emails", [])))
    links = [l for p in ordered for l in p.get("links", [])]
    others = [f"{k}: {v}" for k, v in socials.items() if k not in profs and k not in ("youtube", "tiktok", "instagram")]
    others += [l for l in dict.fromkeys(links) if l not in socials.values() and not any(l == p.get("url") for p in ordered)]
    flags = list(dict.fromkeys(f for mm in [m, *extra_matches] for f in mm.get("red_flags", [])))
    if any(mm.get("competitor_sponsor") for mm in [m, *extra_matches]):
        flags.insert(0, "Sponsored by a competitor")
    windows = [f"{PLATFORMS[p['platform']]}: {p['views_window']}" for p in ordered if p.get("views_window")]
    trend, vs = primary.get("views_trend"), primary.get("engagement_vs_typical")
    last = [p["days_since_last_post"] for p in ordered if p.get("days_since_last_post") is not None]
    agency = (partner or {}).get("agency") or any(agency_hint(p) for p in ordered)
    weeks = (partner or {}).get("weeks") or []

    def url(profile, network):  # a profile we fetched, else the link they gave
        return (profile or {}).get("url") or socials.get(network, "")

    return {
        "Creator key": (partner or {}).get("name") or _short_name(primary.get("name")),
        "Market": country,
        "Country": MARKETS.get(country, {}).get("name", country),
        "Creator / channel": " / ".join(names),
        "Agency": "Yes" if agency else "",
        "Year-week": weeks[-1] if weeks else "",
        "Platform": " + ".join(_platform_label(p) for p in ordered if p["platform"] in ("youtube", "tiktok"))
                    + (" + Instagram" if ig and (yt or tt) else "Instagram" if ig else ""),
        "Niche / content": _niche(m),
        "YT subscribers": (yt or {}).get("followers"),
        "YT views / video": (yt or {}).get("avg_views"),
        "TikTok followers": (tt or {}).get("followers"),
        "TikTok views / video": (tt or {}).get("avg_views"),
        "Instagram followers": (ig or {}).get("followers"),
        "Email": "; ".join(emails),
        "YouTube": url(yt, "youtube"),
        "TikTok": url(tt, "tiktok"),
        "Instagram": url(ig, "instagram"),
        "Other links": "; ".join(others[:6]),
        "Match score": max([m["score"]] + [mm["score"] for mm in extra_matches]),
        "Fit": m.get("fit"),
        "Audience quality": m.get("quality"),
        "Confidence": (m.get("confidence") or {}).get("level", ""),
        "Authenticity": (primary.get("authenticity") or {}).get("score"),
        "Est. price per post (EUR)": f"{price['low']}-{price['high']}" if (price := primary.get("price")) else "",
        "Views counted over": "; ".join(windows),
        "Views trend": f"{trend:+.0%}" if trend is not None else "",
        "Engagement vs typical": vs,
        "Last post (days ago)": min(last) if last else None,
        "Risks / red flags": "; ".join(flags),
        "Why they fit": " | ".join(m.get("why", [])),
        "Verdict": m.get("verdict") or "",
        "Summary": m.get("summary") or "",
        "Status": m.get("status") or "",
        "Past collaborations": ", ".join(weeks) or ("yes" if partner else ""),
    }


def build_rows(company: dict, rows: list[tuple[dict, dict]], creators: dict[str, dict],
               matches: dict[str, dict]) -> list[dict]:
    """One dict per person, in the order of `rows` (the grid's sort)."""
    group = linking.groups(creators)
    ranked = set(matches)
    idx = partners.index(company)
    seen, out = set(), []
    for c, m in rows:
        ids = group.get(c["id"], [c["id"]])
        person = min(ids)
        if person in seen:
            continue
        seen.add(person)
        profs = linking.profiles(ids, creators, prefer={c["id"]} | (ranked & set(ids)))
        profs[c["platform"]] = c
        extra = [matches[i] for i in ids if i != c["id"] and i in matches]
        partner = next((p for p in (partners.find(idx, prof, m) for prof in profs.values()) if p), None)
        out.append(_row(c, m, profs, partner, extra))
    return out


def to_csv(people: list[dict]) -> str:
    buf = io.StringIO()
    w = csv.writer(buf)
    w.writerow([name for name, _ in COLUMNS])
    for p in people:
        w.writerow(["" if p[name] is None else p[name] for name, _ in COLUMNS])
    return "﻿" + buf.getvalue()  # BOM so Excel opens UTF-8 (ä, ö, ß) correctly


def to_xlsx(people: list[dict]) -> bytes:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill

    wb = Workbook()
    ws = wb.active
    ws.title = "Creators"
    ws.append([name for name, _ in COLUMNS])
    for i, cell in enumerate(ws[1]):
        tracker = i < len(TRACKER_COLUMNS)
        cell.font = Font(bold=True, color="FFFFFF" if tracker else "16161A")
        cell.fill = PatternFill("solid", fgColor="16161A" if tracker else "D9D9DE")
        cell.alignment = Alignment(vertical="center", wrap_text=True)
    for p in people:
        ws.append([p[name] for name, _ in COLUMNS])
    for col, (name, width) in enumerate(COLUMNS, start=1):
        letter = ws.cell(row=1, column=col).column_letter
        ws.column_dimensions[letter].width = width
        for cell in ws[letter][1:]:
            if cell.value in (None, ""):
                continue
            if name in NUMBER_COLS:
                cell.number_format = "#,##0"
            elif name in VIEWS_COLS:
                cell.number_format = VIEWS_FORMAT
            elif name in URL_COLS:
                cell.hyperlink = cell.value
                cell.font = Font(color="1D4ED8", underline="single")
    ws.freeze_panes = "B2"
    ws.auto_filter.ref = ws.dimensions
    ws.row_dimensions[1].height = 32

    about = wb.create_sheet("How to read")
    about.append(["Column", "Meaning"])
    for cell in about[1]:
        cell.font = Font(bold=True)
    for row in ABOUT:
        about.append(list(row))
    about.column_dimensions["A"].width = 22
    about.column_dimensions["B"].width = 110
    for row in about.iter_rows(min_row=2):
        row[1].alignment = Alignment(wrap_text=True, vertical="top")
        row[0].alignment = Alignment(vertical="top")
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


# --- The company's own tracker, handed back completed ------------------------------------------------

FILL_COLUMNS = ("Market", "Country", "Creator / channel", "Platform", "Niche / content",
                "YT subscribers", "YT views / video", "TikTok followers", "TikTok views / video")
NOW_COLUMNS = [("YT subscribers now", "YT subscribers"), ("YT views / video now", "YT views / video"),
               ("TikTok followers now", "TikTok followers"), ("TikTok views / video now", "TikTok views / video")]
TRACKER_EXTRA = [("YouTube", 34), ("TikTok", 34), ("Email", 30), ("YT subscribers now", 12), ("YT views / video now", 12),
                 ("TikTok followers now", 12), ("TikTok views / video now", 12), ("Fit", 7), ("Audience quality", 9),
                 ("Filled by Scout", 30), ("Scout lookup", 34)]
LOOKUP_TEXT = {
    "found": "Found", "not_found": "Not found on YouTube or TikTok",
    "twitch": "Only on Twitch: not searched", "no_source": "Not looked up (platform not set up)",
}
TRACKER_ABOUT = [
    ("Your columns", "Your tracker as you uploaded it, row for row. Cells you left empty are filled in where Scout "
                     "found the creator; those cells are highlighted. Nothing you typed is changed."),
    ("... now", "Scout's current numbers, next to what your sheet says. Views: average per video over the last "
                "30 days (90 for less active creators)."),
    ("Fit / Audience quality", "Scout's scores for the creator (0-100), the same as in the app."),
    ("Scout lookup", "Found = a profile with the name from your sheet and a size close to your numbers. "
                     "'Found by name only' = your sheet had no numbers to compare: please check the link."),
]


def _people_for_partners(company: dict, creators: dict[str, dict], matches: dict[str, dict]) -> dict[str, list[str]]:
    """partner name -> creator ids: looked up from the tracker, or found in the results by name."""
    lookup = ((company.get("partners") or {}).get("lookup") or {}).get("results") or {}
    out = {name: [i for i in r.get("ids", []) if i in creators] for name, r in lookup.items()}
    idx = partners.index(company)
    for cid, m in matches.items():
        c = creators.get(cid)
        p = partners.find(idx, c, m) if c else None
        if p and cid not in out.setdefault(p["name"], []):
            out[p["name"]].append(cid)
    return out


def _lookup_note(company: dict, name: str, found: bool) -> str:
    r = (((company.get("partners") or {}).get("lookup") or {}).get("results") or {}).get(name)
    if r:
        if r["status"] == "found":
            how = f" ({r['how']})" if r.get("how") else ""
            return ("Found" if r.get("sure") else "Found by name only: please check") + how
        return LOOKUP_TEXT.get(r["status"], r["status"])
    return "Found by Scout's searches" if found else "Not looked up yet"


def tracker_xlsx(company: dict, creators: dict[str, dict], matches: dict[str, dict]) -> bytes:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill

    info = company.get("partners") or {}
    items = {partners.norm(p["name"]): p for p in info.get("items", [])}
    sheet = info.get("sheet")
    if not sheet:  # imported before sheets were kept: one row per creator from what we parsed
        sheet = {"header": [n for n, _ in TRACKER_COLUMNS], "rows": [
            [p["name"], p["market"], MARKETS.get(p["market"], {}).get("name", ""), " / ".join(p.get("channels", [])[1:]) or p["name"],
             "Yes" if p["agency"] else "", (p["weeks"] or [""])[-1], " + ".join(p["platforms"]), p["niche"],
             p["yt_subs"], None, p["tt_followers"], None] for p in info.get("items", [])]}
    header = sheet["header"]
    col = {h.strip().lower(): i for i, h in enumerate(header)}
    key_i = next((col[h] for h in partners.HEADERS["key"] if h in col), None)
    chan_i = next((col[h] for h in partners.HEADERS["channel"] if h in col), None)
    group = linking.groups(creators)
    people = _people_for_partners(company, creators, matches)
    cache: dict[str, dict | None] = {}

    def person(p: dict) -> dict | None:
        """The row Scout would write for this partner (profiles merged), or None if we don't have them."""
        if p["name"] not in cache:
            ids = [j for i in people.get(p["name"], []) for j in group.get(i, [i])]
            ids = list(dict.fromkeys(ids))
            if not ids:
                cache[p["name"]] = None
            else:
                ranked = [i for i in ids if i in matches]
                primary_id = max(ranked, key=lambda i: matches[i]["score"]) if ranked else ids[0]
                profs = linking.profiles(ids, creators, prefer=set(ranked))
                primary = creators[primary_id]
                profs[primary["platform"]] = primary
                m = matches.get(primary_id) or {"score": None}
                extra = [matches[i] for i in ranked if i != primary_id]
                cache[p["name"]] = _row(primary, m, profs, p, extra)
        return cache[p["name"]]

    wb = Workbook()
    ws = wb.active
    ws.title = "Tracker"
    ws.append(header + [n for n, _ in TRACKER_EXTRA])
    for i, cell in enumerate(ws[1]):
        own = i < len(header)
        cell.font = Font(bold=True, color="FFFFFF" if own else "16161A")
        cell.fill = PatternFill("solid", fgColor="16161A" if own else "D9D9DE")
        cell.alignment = Alignment(vertical="center", wrap_text=True)
    filled_fill = PatternFill("solid", fgColor="FFE3CC")
    for raw in sheet["rows"]:
        row = list(raw) + [None] * (len(header) - len(raw))
        name = (row[key_i] if key_i is not None else None) or (row[chan_i] if chan_i is not None else None)
        p = items.get(partners.norm(name)) if name else None
        got = person(p) if p else None
        filled = []
        if got:
            for column in FILL_COLUMNS:
                i = col.get(column.lower())
                if i is not None and row[i] in (None, "") and got.get(column) not in (None, ""):
                    row[i] = got[column]
                    filled.append(column)
        extra = [(got or {}).get("YouTube"), (got or {}).get("TikTok"), (got or {}).get("Email")] \
            + [(got or {}).get(src) for _, src in NOW_COLUMNS] \
            + [(got or {}).get("Fit"), (got or {}).get("Audience quality"), ", ".join(filled),
               _lookup_note(company, p["name"], bool(got)) if p else ""]
        ws.append(row + extra)
        r = ws.max_row
        for column in filled:
            ws.cell(row=r, column=col[column.lower()] + 1).fill = filled_fill
    widths = {n.lower(): w for n, w in TRACKER_COLUMNS}
    for i, h in enumerate(header + [n for n, _ in TRACKER_EXTRA], start=1):
        letter = ws.cell(row=1, column=i).column_letter
        ws.column_dimensions[letter].width = widths.get(h.lower()) or dict(TRACKER_EXTRA).get(h, 14)
        for cell in ws[letter][1:]:
            if not isinstance(cell.value, (int, float)):
                if h in URL_COLS and cell.value:
                    cell.hyperlink = cell.value
                    cell.font = Font(color="1D4ED8", underline="single")
                continue
            if h in NUMBER_COLS or h in ("YT subscribers now", "TikTok followers now"):
                cell.number_format = "#,##0"
            elif h in VIEWS_COLS or h.endswith("views / video now"):
                cell.number_format = VIEWS_FORMAT
    ws.freeze_panes = "B2"
    ws.auto_filter.ref = ws.dimensions
    ws.row_dimensions[1].height = 32
    about = wb.create_sheet("How to read")
    about.append(["Column", "Meaning"])
    for cell in about[1]:
        cell.font = Font(bold=True)
    for line in TRACKER_ABOUT:
        about.append(list(line))
    about.column_dimensions["A"].width = 24
    about.column_dimensions["B"].width = 110
    for row in about.iter_rows(min_row=2):
        row[1].alignment = Alignment(wrap_text=True, vertical="top")
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
